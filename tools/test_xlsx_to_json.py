"""Cobre a sanitizacao de CPF em tools/xlsx_to_json.diretorias(), a
normalizacao de tools/xlsx_to_json.corpo_funcional() e a transcricao da
estrutura remuneratoria a partir do XLSX.

Cada teste monta uma planilha minima em tmp_path e aponta ORIGEM para la:
nada aqui toca downloads/ ou data/, salvo os que leem a planilha oficial.

Rodar com: python -m pytest tools/test_xlsx_to_json.py -q
"""
import json
import pathlib
import sys

import openpyxl
import pytest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

import xlsx_to_json  # noqa: E402

NOTA_INTERNA = "FUG - Administrativo 2 adm2: não tem o CPF na planilha que tenho acesso"
MASCARA_RITA = "564.XXX.XXX-34"


def planilha_diretorias(caminho, abas):
    """abas: dict titulo -> lista de (unidade, cargo, nome, cpf)."""
    wb = openpyxl.Workbook()
    primeira = True
    for titulo, linhas in abas.items():
        ws = wb.active if primeira else wb.create_sheet()
        primeira = False
        ws.title = titulo
        for i, (unidade, cargo, nome, cpf) in enumerate(linhas, start=5):
            ws.cell(i, 1, unidade)
            ws.cell(i, 2, cargo)
            ws.cell(i, 3, nome)
            ws.cell(i, 4, cpf)
    wb.save(caminho)
    return caminho


def _registros_por_colegiado(tmp_path, monkeypatch, abas):
    origem = tmp_path / "downloads"
    origem.mkdir()
    planilha_diretorias(origem / "diretorias-2025.xlsx", abas)
    monkeypatch.setattr(xlsx_to_json, "ORIGEM", origem)
    return {
        c["id"]: c["registros"] for c in xlsx_to_json.diretorias()["colegiados"]
    }


def test_diretorias_sanitiza_cpf_ocupado_sem_marcar_vago(tmp_path, monkeypatch):
    registros = _registros_por_colegiado(
        tmp_path,
        monkeypatch,
        {
            "Conselho Curador": [
                (
                    "Nacional",
                    "VICE-PRESIDENTE",
                    "Katia Damiana Alves Pereira Lobo",
                    "773.XXX.XXX–04",
                ),
                (
                    "Nacional",
                    "SUPLENTE",
                    "Fulano Nota Interna",
                    NOTA_INTERNA,
                ),
            ]
        },
    )["conselho-curador"]

    katia = next(r for r in registros if r["nome"] == "Katia Damiana Alves Pereira Lobo")
    assert katia["cpf"] == "773.XXX.XXX-04"
    assert katia["vago"] is False

    fulano = next(r for r in registros if r["nome"] == "Fulano Nota Interna")
    assert fulano["cpf"] == "—"
    assert fulano["vago"] is False


def test_diretorias_publica_mascara_conhecida_da_maria_rita(tmp_path, monkeypatch):
    colegiados = _registros_por_colegiado(
        tmp_path,
        monkeypatch,
        {
            "DIRETORIAS ESTADUAIS": [
                (
                    "São Paulo",
                    "VICE-PRESIDENTE",
                    "MARIA RITA CARRA NAVARRO",
                    MASCARA_RITA,
                ),
            ],
            "CONSELHO CURADOR": [
                (
                    "São Paulo",
                    "VICE-PRESIDENTE",
                    "MARIA RITA CARRA NAVARRO",
                    NOTA_INTERNA,
                ),
                (
                    "Nacional",
                    "PRESIDENTE",
                    "WELLINGTON MOREIRA FRANCO",
                    "103.XXX.XXX-91",
                ),
            ],
        },
    )

    estaduais = next(
        r
        for r in colegiados["diretorias-estaduais"]
        if r["nome"] == "Maria Rita Carra Navarro"
    )
    nomes_curador = [r["nome"] for r in colegiados["conselho-curador"]]

    assert estaduais["cpf"] == MASCARA_RITA
    assert estaduais["vago"] is False
    assert "Maria Rita Carra Navarro" not in nomes_curador
    assert "Wellington Moreira Franco" in nomes_curador


def test_conselho_curador_mantem_so_unidade_nacional(tmp_path, monkeypatch):
    colegiados = _registros_por_colegiado(
        tmp_path,
        monkeypatch,
        {
            "CONSELHO CURADOR": [
                ("Nacional", "PRESIDENTE", "WELLINGTON MOREIRA FRANCO", "103.XXX.XXX-91"),
                ("nacional", "VICE-PRESIDENTE", "CARLOS ALBERTO CHIODINI", "005.XXX.XXX-42"),
                ("São Paulo", "VICE-PRESIDENTE", "MARIA RITA CARRA NAVARRO", MASCARA_RITA),
                ("Rio de Janeiro", "CONSELHEIRO TITULAR", "ALGUEM ESTADUAL", "111.XXX.XXX-11"),
            ],
            "DIRETORIAS ESTADUAIS": [
                ("São Paulo", "PRESIDENTE", "ALGUEM ESTADUAL", "111.XXX.XXX-11"),
            ],
        },
    )
    curador = colegiados["conselho-curador"]
    assert [r["nome"] for r in curador] == [
        "Wellington Moreira Franco",
        "Carlos Alberto Chiodini",
    ]
    assert all(r["unidade"].casefold() == "nacional" for r in curador)
    assert [r["unidade"] for r in colegiados["diretorias-estaduais"]] == ["São Paulo"]


def test_conselho_curador_oficial_tem_15_nacionais():
    curador = next(
        c for c in xlsx_to_json.diretorias()["colegiados"] if c["id"] == "conselho-curador"
    )
    assert curador["total"] == 15
    assert len(curador["registros"]) == 15
    assert all(r["unidade"].casefold() == "nacional" for r in curador["registros"])
    nomes = [r["nome"] for r in curador["registros"]]
    assert "Wellington Moreira Franco" in nomes
    assert "Maria Rita Carra Navarro" not in nomes



def planilha_corpo_funcional(caminho, linhas):
    """linhas: lista de (estado, nome, cargo). Cabecalho na linha 1; dados a partir da 2."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1, "ESTADOS")
    ws.cell(1, 2, "FUNCIONÁRIO")
    ws.cell(1, 3, "CARGO")
    for i, (estado, nome, cargo) in enumerate(linhas, start=2):
        ws.cell(i, 1, estado)
        ws.cell(i, 2, nome)
        ws.cell(i, 3, cargo)
    wb.save(caminho)
    return caminho


def _corpo_funcional(tmp_path, monkeypatch, linhas):
    origem = tmp_path / "downloads"
    origem.mkdir()
    planilha_corpo_funcional(origem / "corpo-funcional-2025.xlsx", linhas)
    monkeypatch.setattr(xlsx_to_json, "ORIGEM", origem)
    return xlsx_to_json.corpo_funcional()


def test_corpo_funcional_propaga_estados_vazio(tmp_path, monkeypatch):
    dados = _corpo_funcional(
        tmp_path,
        monkeypatch,
        [
            ("RIO GRANDE DO SUL", "LUIZ FERNANDO RODRIGUES DA ROSA", "Auxiliar"),
            ("", "Sanmartin Ciceri", "Estagiário"),
        ],
    )
    por_nome = {r["nome"]: r for r in dados["registros"]}
    assert por_nome["Sanmartin Ciceri"]["unidade"] == "Rio Grande do Sul"


def test_corpo_funcional_aplica_correcoes_nomeadas(tmp_path, monkeypatch):
    dados = _corpo_funcional(
        tmp_path,
        monkeypatch,
        [
            ("RORAIMA RR", "Elaíne Santos deJesus", "Servilços Gerais"),
            ("", "FABIOLA DOS SANTOS LIMA", "Auxiliar Administrativo Júnior"),
            ("DISTRITO FEDERAL", "ANA SILVA", "Auxiliar Administrativo Senior"),
        ],
    )
    por_nome = {r["nome"]: r for r in dados["registros"]}

    elaine = por_nome["Elaíne Santos de Jesus"]
    assert elaine["cargo"] == "Auxiliar de Serviços Gerais"
    assert elaine["unidade"] == "Roraima"

    fabiola = por_nome["Fabiola dos Santos Lima"]
    assert fabiola["unidade"] == "Roraima"

    assert por_nome["Ana Silva"]["cargo"] == "Auxiliar Administrativo Sênior"

    publicado = json.dumps(dados, ensure_ascii=False)
    assert "Roraima Rr" not in publicado
    assert "Servilços" not in publicado
    assert "deJesus" not in publicado
    assert "Senior" not in publicado


def test_corpo_funcional_ordena_por_nome(tmp_path, monkeypatch):
    dados = _corpo_funcional(
        tmp_path,
        monkeypatch,
        [
            ("ACRE", "ZECA ULTIMO", "Copeiro"),
            ("CEARÁ", "ANA PRIMEIRA", "Motorista"),
        ],
    )
    assert [r["nome"] for r in dados["registros"]] == ["Ana Primeira", "Zeca Ultimo"]


def test_corpo_funcional_nao_traz_cpf_nem_valor_monetario(tmp_path, monkeypatch):
    dados = _corpo_funcional(
        tmp_path,
        monkeypatch,
        [("DISTRITO FEDERAL", "JOÃO DA SILVA", "Copeiro")],
    )
    publicado = json.dumps(dados, ensure_ascii=False)
    assert "cpf" not in publicado.lower()
    assert "R$" not in publicado
    for registro in dados["registros"]:
        assert set(registro) == {"unidade", "nome", "cargo"}


def test_corpo_funcional_transcreve_29_registros_da_planilha_oficial():
    dados = xlsx_to_json.corpo_funcional()
    assert "ano" in dados
    assert dados["total"] == 29
    assert len(dados["registros"]) == 29
    assert dados["total"] == len(dados["registros"])


def test_diretorias_nome_vazio_publica_vacancia(tmp_path, monkeypatch):
    registros = _registros_por_colegiado(
        tmp_path,
        monkeypatch,
        {"Diretorias Estaduais": [("Ceará", "PRESIDENTE", "", "")]},
    )["diretorias-estaduais"]

    assert len(registros) == 1
    assert registros[0]["nome"] == "VACÂNCIA"
    assert registros[0]["vago"] is True
    assert registros[0]["cpf"] == "—"
    assert xlsx_to_json.CARGO_VAGO == "VACÂNCIA"


CARGOS_ER_2026 = [
    ("Secretário Executivo", 10, 1, 1, "R$ 25.000,00", "R$ 35.000,00", "R$ 50.000,00"),
    ("Secretário Executivo Adjunto", 10, 1, 1, "R$ 20.000,00", "R$ 30.000,00", "R$ 35.000,00"),
    ("Procurador Jurídico", 9, 1, 1, "R$ 14.000,00", "R$ 17.000,00", "R$ 20.000,00"),
    ("Gerente", 8, 3, 0, "R$ 14.000,00", "R$ 17.000,00", "R$ 20.000,00"),
    ("Chefe de Gabinete Diretoria Executiva", 7, 1, 1, "R$ 12.511,20", "R$ 15.000,00", "R$ 18.000,00"),
    ("Supervisor", 6, 10, 6, "R$ 11.468,60", "R$ 13.000,00", "R$ 16.000,00"),
    ("Coordenador", 5, 10, 6, "R$ 9.195,74", "R$ 11.000,00", "R$ 14.000,00"),
    ("Assistente Administrativo I", 4, 35, 0, "R$ 5.200,00", "R$ 6.234,75", "R$ 6.540,00"),
    ("Assistente Administrativo II", 4, 35, 0, "R$ 6.828,39", "R$ 6.900,00", "R$ 7.280,00"),
    ("Assistente Administrativo III", 4, 35, 1, "R$ 7.595,68", "R$ 7.900,00", "R$ 8.500,00"),
    ("Auxiliar Administrativo Junior", 3, 35, 6, "R$ 3.239,03", "R$ 3.500,00", "R$ 3.790,00"),
    ("Auxiliar Administrativo Pleno", 3, 35, 1, "R$ 3.800,00", "R$ 4.000,00", "R$ 4.400,00"),
    ("Auxiliar Administrativo Sênior", 3, 35, 3, "R$ 4.671,60", "R$ 4.900,00", "R$ 5.100,00"),
    ("Trainee", 3, 5, 0, "R$ 1.848,00", "—", "—"),
    ("Estagiário Ensino Médio", 2, 5, 2, "R$ 1.300,00", "—", "—"),
    ("Estagiário Superior", 2, 5, 2, "R$ 1.500,00", "—", "—"),
    ("Motorista", 2, 2, 2, "R$ 5.535,83", "R$ 5.900,00", "R$ 6.500,00"),
    ("Aux. Serv. Gerais", 1, 2, 1, "R$ 1.973,58", "R$ 2.100,00", "R$ 3.000,00"),
    ("Copeiro", 1, 1, 1, "R$ 3.704,40", "R$ 4.100,00", "R$ 5.947,41"),
    ("Jovem Aprendiz", 1, 5, 3, "R$ 713,00", "—", "—"),
]
CAMPOS_CARGO = {
    "cargo",
    "nivel",
    "quadro_total",
    "quadro_preenchido",
    "salario_base",
    "faixa_media",
    "faixa_maxima",
    "adicionais",
}
CARGOS_QP_UM = (
    "Secretário Executivo",
    "Secretário Executivo Adjunto",
    "Procurador Jurídico",
)
CARGOS_QP_ZERO = (
    "Gerente",
    "Assistente Administrativo I",
    "Assistente Administrativo II",
    "Trainee",
)
CARGOS_SEM_FAIXA = (
    "Trainee",
    "Estagiário Ensino Médio",
    "Estagiário Superior",
    "Jovem Aprendiz",
)


def planilha_remuneratoria(caminho, linhas, cabecalho_salario="Salário Base 2026"):
    """linhas: (qt, qp, cargo, nivel, base, media, maxima). Cabecalho na linha 1."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1, "Q. T.")
    ws.cell(1, 2, "Q. P.")
    ws.cell(1, 3, "Cargo")
    ws.cell(1, 4, "Nível")
    ws.cell(1, 5, cabecalho_salario)
    ws.cell(1, 6, "Faixa Média")
    ws.cell(1, 7, "Faixa Máxima")
    for i, (qt, qp, cargo, nivel, base, media, maxima) in enumerate(linhas, start=2):
        ws.cell(i, 1, qt)
        ws.cell(i, 2, qp)
        ws.cell(i, 3, cargo)
        ws.cell(i, 4, nivel)
        ws.cell(i, 5, base)
        ws.cell(i, 6, media)
        ws.cell(i, 7, maxima)
    wb.save(caminho)
    return caminho


def _estrutura_remuneratoria(tmp_path, monkeypatch, linhas, cabecalho_salario="Salário Base 2026"):
    origem = tmp_path / "downloads"
    origem.mkdir()
    planilha_remuneratoria(
        origem / "estrutura-remuneratoria-2026.xlsx",
        linhas,
        cabecalho_salario=cabecalho_salario,
    )
    monkeypatch.setattr(xlsx_to_json, "ORIGEM", origem)
    return xlsx_to_json.estrutura_remuneratoria()


def test_estrutura_remuneratoria_formata_tipos_mistos_e_travessao(tmp_path, monkeypatch):
    dados = _estrutura_remuneratoria(
        tmp_path,
        monkeypatch,
        [
            (1, 0, "Secretário Executivo", 10, "R$ 25.000,00", "R$ 35.000,00", 50000.0),
            (35, 0, "Assistente Administrativo II", 4, "R$ 6.828,39", "R$6.900,00", "R$ 7.280,00"),
            (5, 0, "Trainee", 3, 1848.0, "R$ –", "R$ -"),
            (5, 3, "Jovem Aprendiz ", 1, "R$ 713,00", "R$ –", "R$ –"),
        ],
    )
    por_cargo = {c["cargo"]: c for c in dados["cargos"]}

    assert dados["ano"] == 2026
    assert por_cargo["Secretário Executivo"]["faixa_maxima"] == "R$ 50.000,00"
    assert por_cargo["Assistente Administrativo II"]["faixa_media"] == "R$ 6.900,00"
    assert por_cargo["Trainee"]["salario_base"] == "R$ 1.848,00"
    assert por_cargo["Trainee"]["faixa_media"] == "—"
    assert por_cargo["Trainee"]["faixa_maxima"] == "—"
    assert "Jovem Aprendiz" in por_cargo
    assert "Jovem Aprendiz " not in por_cargo
    zeros = ("0", "R$ 0", "R$ 0,00")
    for cargo in dados["cargos"]:
        assert cargo["faixa_media"] not in zeros
        assert cargo["faixa_maxima"] not in zeros
        assert cargo["adicionais"] == "—"


def test_estrutura_remuneratoria_transcreve_qp_sem_override(tmp_path, monkeypatch):
    dados = _estrutura_remuneratoria(
        tmp_path,
        monkeypatch,
        [
            (1, 0, "Secretário Executivo", 10, 25000, 35000, 50000),
            (3, 0, "Gerente", 8, 14000, 17000, 20000),
        ],
    )
    por_cargo = {c["cargo"]: c for c in dados["cargos"]}
    assert por_cargo["Secretário Executivo"]["quadro_preenchido"] == 0
    assert por_cargo["Gerente"]["quadro_preenchido"] == 0


def test_estrutura_remuneratoria_ordena_por_nivel_depois_cargo(tmp_path, monkeypatch):
    dados = _estrutura_remuneratoria(
        tmp_path,
        monkeypatch,
        [
            (1, 1, "Copeiro", 1, "R$ 3.704,40", "R$ 4.100,00", "R$ 5.947,41"),
            (1, 1, "Aux. Serv. Gerais", 1, "R$ 1.973,58", "R$ 2.100,00", "R$ 3.000,00"),
            (1, 0, "Secretário Executivo Adjunto", 10, 20000, "R$ 30.000,00", "R$ 35.000,00"),
            (1, 0, "Secretário Executivo", 10, "R$ 25.000,00", "R$ 35.000,00", 50000),
        ],
    )
    assert [c["cargo"] for c in dados["cargos"]] == [
        "Secretário Executivo",
        "Secretário Executivo Adjunto",
        "Aux. Serv. Gerais",
        "Copeiro",
    ]


def test_estrutura_remuneratoria_ausente_nao_escreve_json(tmp_path, monkeypatch):
    origem = tmp_path / "downloads"
    origem.mkdir()
    destino = tmp_path / "data"
    destino.mkdir()
    monkeypatch.setattr(xlsx_to_json, "ORIGEM", origem)
    monkeypatch.setattr(xlsx_to_json, "DESTINO", destino)
    monkeypatch.setattr(
        xlsx_to_json, "corpo_funcional", lambda: {"ano": 2025, "total": 0, "registros": []}
    )
    monkeypatch.setattr(
        xlsx_to_json, "diretorias", lambda: {"ano": 2025, "colegiados": []}
    )

    with pytest.raises(FileNotFoundError):
        xlsx_to_json.estrutura_remuneratoria()
    assert xlsx_to_json.main() == 1
    assert not (destino / "estrutura-remuneratoria.json").exists()


def test_estrutura_remuneratoria_transcreve_20_cargos_do_xlsx_oficial():
    dados = xlsx_to_json.estrutura_remuneratoria()
    por_cargo = {c["cargo"]: c for c in dados["cargos"]}

    assert dados["ano"] == 2026
    assert set(dados) == {"ano", "total_quadro_preenchido", "cargos"}
    assert len(dados["cargos"]) == 20
    assert [c["cargo"] for c in dados["cargos"]] == [linha[0] for linha in CARGOS_ER_2026]
    for extraido, esperado in zip(dados["cargos"], CARGOS_ER_2026, strict=True):
        cargo, nivel, qt, qp, base, media, maxima = esperado
        assert set(extraido) == CAMPOS_CARGO
        assert extraido == {
            "cargo": cargo,
            "nivel": nivel,
            "quadro_total": qt,
            "quadro_preenchido": qp,
            "salario_base": base,
            "faixa_media": media,
            "faixa_maxima": maxima,
            "adicionais": "—",
        }

    assert por_cargo["Secretário Executivo"]["faixa_maxima"] == "R$ 50.000,00"
    assert por_cargo["Chefe de Gabinete Diretoria Executiva"]["salario_base"] == "R$ 12.511,20"
    assert por_cargo["Estagiário Superior"]["salario_base"] == "R$ 1.500,00"
    assert por_cargo["Copeiro"]["faixa_maxima"] == "R$ 5.947,41"
    for nome in CARGOS_QP_UM:
        assert por_cargo[nome]["quadro_preenchido"] == 1
    for nome in CARGOS_QP_ZERO:
        assert por_cargo[nome]["quadro_preenchido"] == 0
    for nome in CARGOS_SEM_FAIXA:
        assert por_cargo[nome]["faixa_media"] == "—"
        assert por_cargo[nome]["faixa_maxima"] == "—"
    assert dados["total_quadro_preenchido"] == sum(
        c["quadro_preenchido"] for c in dados["cargos"]
    )
