"""Converte as planilhas oficiais da FUG para os JSON normalizados em data/.

Fonte de verdade: os arquivos .xlsx em downloads/. Reexecute este script
sempre que as planilhas forem atualizadas e depois rode tools/build_tables.py.
"""
import json
import pathlib
import re
import sys
import unicodedata

import openpyxl

RAIZ = pathlib.Path(__file__).resolve().parent.parent
ORIGEM = RAIZ / "downloads"
DESTINO = RAIZ / "data"
ANO = 2025

# Particulas que permanecem em minusculas dentro de nomes proprios.
PARTICULAS = {"de", "da", "do", "das", "dos", "e", "di", "du", "del", "van", "von", "la"}

# Correcoes de digitacao encontradas nas planilhas de origem, aplicadas na
# normalizacao para que a pagina publicada nao reproduza o erro.
CORRECOES_CARGO = {
    "Servilços Gerais": "Auxiliar de Serviços Gerais",
    "Auxiliar Administrativo Senior": "Auxiliar Administrativo Sênior",
}

CORRECOES_NOME = {
    "Elaíne Santos deJesus": "Elaíne Santos de Jesus",
}

CORRECOES_UNIDADE = {
    "RORAIMA RR": "Roraima",
}

# CPF mascarado conhecido quando a celula da planilha traz nota interna
# em vez da mascara. So se aplica se sanitizar_cpf cair em "—" — mascara
# valida na propria linha nao e sobrescrita.
CORRECOES_CPF = {
    "Maria Rita Carra Navarro": "564.XXX.XXX-34",
}

# Exibido quando a planilha traz o cargo sem titular associado.
CARGO_VAGO = "VACÂNCIA"
ARQUIVO_REMUNERATORIA = "estrutura-remuneratoria-2026.xlsx"

# CPF ocupado publicado: mascara da planilha, com hifen ASCII. Qualquer outro
# valor (nota interna, vazio, lixo) vira o mesmo traco dos cargos vagos.
_MASCARA_CPF = re.compile(r"^\d{3}\.XXX\.XXX-\d{2}$")


def limpar(valor):
    if valor is None:
        return ""
    texto = str(valor)
    # Remove hifen suave e demais caracteres de formatacao invisiveis que vieram
    # de colagem entre editores e quebram a busca no navegador.
    texto = "".join(c for c in texto if unicodedata.category(c) != "Cf")
    return re.sub(r"\s+", " ", texto).strip()


def nome_proprio(texto):
    if not texto or texto != texto.upper() and texto != texto.lower():
        return texto

    def palavra(match):
        p = match.group(0)
        if p.lower() in PARTICULAS:
            return p.lower()
        return p[0].upper() + p[1:].lower()

    convertido = re.sub(r"[^\s\-/()]+", palavra, texto.lower())
    # Siglas de UF entre parenteses voltam para maiusculas: "(rs)" -> "(RS)".
    convertido = re.sub(r"\((\w{2})\)", lambda m: f"({m.group(1).upper()})", convertido)
    return convertido[0].upper() + convertido[1:] if convertido else convertido


def sanitizar_cpf(cpf):
    """Hifen Unicode vira ASCII; ocupado fora de ddd.XXX.XXX-dd vira —."""
    texto = ("" if cpf is None else str(cpf)).replace("\u2013", "-")
    if _MASCARA_CPF.fullmatch(texto):
        return texto
    return "—"


def escrever(nome, conteudo):
    caminho = DESTINO / nome
    caminho.write_text(
        json.dumps(conteudo, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    visivel = caminho
    try:
        visivel = caminho.relative_to(RAIZ)
    except ValueError:
        pass
    print(f"  {visivel}")


def corpo_funcional():
    ws = openpyxl.load_workbook(ORIGEM / "corpo-funcional-2025.xlsx", data_only=True).active
    registros = []
    estado_atual = ""
    for linha in ws.iter_rows(min_row=2, values_only=True):
        estado, nome, cargo = (limpar(c) for c in linha[:3])
        if not nome:
            continue
        estado_atual = estado or estado_atual
        unidade = CORRECOES_UNIDADE.get(estado_atual, estado_atual)
        cargo = CORRECOES_CARGO.get(cargo, cargo)
        nome = CORRECOES_NOME.get(nome, nome)
        registros.append(
            {
                "unidade": nome_proprio(unidade),
                "nome": nome_proprio(nome),
                "cargo": cargo,
            }
        )
    registros.sort(key=lambda r: r["nome"])
    return {"ano": ANO, "total": len(registros), "registros": registros}


def diretorias():
    wb = openpyxl.load_workbook(ORIGEM / "diretorias-2025.xlsx", data_only=True)
    colegiados = []
    for ws in wb.worksheets:
        registros = []
        for linha in ws.iter_rows(min_row=5, values_only=True):
            unidade, cargo, nome, cpf = (limpar(c) for c in linha[:4])
            if not cargo:
                continue
            nome_pub = nome_proprio(nome) if nome else CARGO_VAGO
            cpf_pub = sanitizar_cpf(cpf) if nome else "—"
            if nome and cpf_pub == "—":
                cpf_pub = sanitizar_cpf(CORRECOES_CPF.get(nome_pub, cpf))
            registros.append(
                {
                    "unidade": unidade,
                    "cargo": cargo,
                    "nome": nome_pub,
                    "cpf": cpf_pub,
                    "vago": not nome,
                }
            )
        id_colegiado = re.sub(r"[^a-z0-9]+", "-", nome_proprio(ws.title).lower()).strip("-")
        # A aba mistura o colegiado nacional com linhas de UF coladas por engano.
        if id_colegiado == "conselho-curador":
            registros = [r for r in registros if r["unidade"].casefold() == "nacional"]
        if not registros:
            continue
        colegiados.append(
            {
                "id": id_colegiado,
                "nome": nome_proprio(ws.title),
                "total": len(registros),
                "registros": registros,
            }
        )
    return {"ano": ANO, "colegiados": colegiados}


def _inteiro(valor):
    if valor is None or valor == "":
        return 0
    return int(valor)


def valor_monetario(celula):
    """Numero vira R$ X.XXX,00; R$ – / R$ - viram —; R$ colado ganha espaço."""
    if celula is None or celula == "":
        return "—"
    if isinstance(celula, (int, float)):
        texto = f"{float(celula):,.2f}"
        return "R$ " + texto.replace(",", "X").replace(".", ",").replace("X", ".")
    texto = limpar(celula)
    if re.fullmatch(r"R\$\s*[–\-]", texto):
        return "—"
    return re.sub(r"^R\$\s*", "R$ ", texto)


def estrutura_remuneratoria():
    caminho = ORIGEM / ARQUIVO_REMUNERATORIA
    if not caminho.exists():
        raise FileNotFoundError(
            f"Planilha ausente: {caminho}. "
            "Nada foi escrito em data/estrutura-remuneratoria.json."
        )
    ws = openpyxl.load_workbook(caminho, data_only=True).active
    cabecalho = limpar(ws.cell(1, 5).value)
    ano_m = re.search(r"(20\d{2})", cabecalho)
    ano = int(ano_m.group(1)) if ano_m else 2026
    cargos = []
    for linha in ws.iter_rows(min_row=2, max_col=7, values_only=True):
        qt, qp, cargo, nivel, base, media, maxima = linha[:7]
        cargo = limpar(cargo)
        if not cargo:
            continue
        cargos.append(
            {
                "cargo": cargo,
                "nivel": _inteiro(nivel),
                "quadro_total": _inteiro(qt),
                "quadro_preenchido": _inteiro(qp),
                "salario_base": valor_monetario(base),
                "faixa_media": valor_monetario(media),
                "faixa_maxima": valor_monetario(maxima),
                "adicionais": "—",
            }
        )
    cargos.sort(key=lambda c: (-c["nivel"], c["cargo"]))
    return {
        "ano": ano,
        "total_quadro_preenchido": sum(c["quadro_preenchido"] for c in cargos),
        "cargos": cargos,
    }


def main():
    DESTINO.mkdir(exist_ok=True)
    print("Gerando JSON a partir das planilhas oficiais:")
    try:
        remuneratoria = estrutura_remuneratoria()
    except FileNotFoundError as erro:
        print(f"erro: {erro}", file=sys.stderr)
        return 1
    escrever("corpo-funcional.json", corpo_funcional())
    escrever("diretorias.json", diretorias())
    escrever("estrutura-remuneratoria.json", remuneratoria)
    return 0


if __name__ == "__main__":
    sys.exit(main())
